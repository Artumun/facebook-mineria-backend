import os
import re
import io
import time
import json
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
import requests
import pymysql
from dotenv import load_dotenv
from flask import Flask, request, jsonify, send_file, Response
from flask_cors import CORS
from werkzeug.security import generate_password_hash, check_password_hash
from playwright.sync_api import sync_playwright
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import PieChart, Reference

RUTA_ENV = os.path.join(os.path.dirname(os.path.abspath(__file__)), '.env')
load_dotenv(dotenv_path=RUTA_ENV, override=True)

app = Flask(__name__)
CORS(app)

ARCHIVO_SESION_FB = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fb_sesion_activa.json')

GEMINI_API_KEY = (
    os.getenv('GEMINI_API_KEY')
    or os.getenv('API_KEY')
    or os.getenv('GOOGLE_API_KEY')
    or ''
).strip().strip("'").strip('"')

if len(GEMINI_API_KEY) > 10:
    print(f"[Motor IA] Gemini 2.5 Flash configurado con clave: {GEMINI_API_KEY[:8]}...")
else:
    print("[Motor IA] ERROR: No se encontró GEMINI_API_KEY en el archivo .env")

SCRAPER_EXECUTOR = ThreadPoolExecutor(max_workers=1)
NAVEGADOR_RAM = {
    'pw': None,
    'browser': None,
    'context': None,
    'autenticado': False
}

ESTADO_PROGRESO = {
    'activo': False,
    'porcentaje': 0,
    'techo_etapa': 15,
    'etapa': 'En espera de consulta...',
    'tiempo_transcurrido': 0,
    'tiempo_restante': 25,
    'inicio_ts': 0.0,
    'meta_holgada_ts': 0.0
}

def iniciar_medidor_progreso(colchon_segundos=25):
    ahora = time.time()
    ESTADO_PROGRESO['activo'] = True
    ESTADO_PROGRESO['porcentaje'] = 8
    ESTADO_PROGRESO['techo_etapa'] = 28
    ESTADO_PROGRESO['etapa'] = 'Conectando con la publicación en Facebook...'
    ESTADO_PROGRESO['tiempo_transcurrido'] = 0
    ESTADO_PROGRESO['tiempo_restante'] = colchon_segundos
    ESTADO_PROGRESO['inicio_ts'] = ahora
    ESTADO_PROGRESO['meta_holgada_ts'] = ahora + colchon_segundos

def actualizar_progreso(porcentaje, etapa, techo_etapa=None, segundos_holgados=None):
    ahora = time.time()
    transcurrido = max(0, int(ahora - ESTADO_PROGRESO['inicio_ts']))
    ESTADO_PROGRESO['tiempo_transcurrido'] = transcurrido
    ESTADO_PROGRESO['porcentaje'] = max(ESTADO_PROGRESO['porcentaje'], int(min(100, porcentaje)))
    ESTADO_PROGRESO['etapa'] = etapa

    if techo_etapa is not None:
        ESTADO_PROGRESO['techo_etapa'] = max(ESTADO_PROGRESO['porcentaje'], int(min(99, techo_etapa)))

    if porcentaje >= 100:
        ESTADO_PROGRESO['activo'] = False
        ESTADO_PROGRESO['porcentaje'] = 100
        ESTADO_PROGRESO['tiempo_restante'] = 0
        return

    if segundos_holgados is not None:
        nuevo_restante = max(4, int(segundos_holgados) + 4)
        ESTADO_PROGRESO['tiempo_restante'] = min(ESTADO_PROGRESO['tiempo_restante'], nuevo_restante)
        ESTADO_PROGRESO['meta_holgada_ts'] = ahora + ESTADO_PROGRESO['tiempo_restante']

def get_db_connection():
    return pymysql.connect(
        host=os.getenv('DB_HOST', 'localhost'),
        user=os.getenv('DB_USER', 'root'),
        password=os.getenv('DB_PASSWORD', ''),
        database=os.getenv('DB_NAME', 'mineria_facebook_db'),
        cursorclass=pymysql.cursors.DictCursor
    )

def asegurar_columnas_bd():
    columnas_analisis = [
        ("contexto_post", "TEXT"),
        ("contexto_multimedia", "TEXT"),
        ("reacciones_lista_json", "TEXT"),
        ("total_bots", "INT DEFAULT 0")
    ]
    columnas_comentarios = [
        ("es_bot", "TINYINT(1) DEFAULT 0"),
        ("motivo_bot", "VARCHAR(120) DEFAULT 'Orgánico'"),
        ("avatar_url", "TEXT"),
        ("confianza_ia", "FLOAT DEFAULT 0")
    ]
    try:
        conn = get_db_connection()
        with conn.cursor() as cursor:
            for nombre_col, tipo_col in columnas_analisis:
                try:
                    cursor.execute(f"ALTER TABLE analisis ADD COLUMN {nombre_col} {tipo_col}")
                except Exception:
                    pass
            for nombre_col, tipo_col in columnas_comentarios:
                try:
                    cursor.execute(f"ALTER TABLE comentarios ADD COLUMN {nombre_col} {tipo_col}")
                except Exception:
                    pass
            conn.commit()
        conn.close()
    except Exception as e:
        print(f"Aviso al verificar esquema de BD: {e}")

def _iniciar_navegador_invisible_worker():
    pw = sync_playwright().start()
    browser = pw.chromium.launch(
        headless=True,
        args=[
            '--headless=new',
            '--disable-blink-features=AutomationControlled',
            '--disable-gpu',
            '--disable-extensions',
            '--no-sandbox'
        ]
    )

    sesion_valida = False
    if os.path.exists(ARCHIVO_SESION_FB):
        try:
            with open(ARCHIVO_SESION_FB, 'r', encoding='utf-8') as f:
                cookies_guardadas = json.load(f).get('cookies', [])
                if any(c.get('name') == 'c_user' for c in cookies_guardadas):
                    sesion_valida = True
        except Exception:
            sesion_valida = False

    opciones_contexto = {
        'locale': 'es-MX',
        'user_agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36',
        'viewport': {'width': 1366, 'height': 900}
    }

    if sesion_valida:
        opciones_contexto['storage_state'] = ARCHIVO_SESION_FB

    context = browser.new_context(**opciones_contexto)
    context.add_init_script("Object.defineProperty(navigator, 'webdriver', {get: () => undefined})")

    context.route(
        "**/*",
        lambda route: route.abort()
        if route.request.resource_type in ["media", "font"]
        else route.continue_()
    )

    NAVEGADOR_RAM['pw'] = pw
    NAVEGADOR_RAM['browser'] = browser
    NAVEGADOR_RAM['context'] = context
    NAVEGADOR_RAM['autenticado'] = sesion_valida
    print(f"[Motor Web] Navegador invisible listo (Sesión de Facebook activa: {'SÍ' if sesion_valida else 'NO'}).")

asegurar_columnas_bd()
SCRAPER_EXECUTOR.submit(_iniciar_navegador_invisible_worker).result()

def traducir_nombre_reaccion(identificador):
    t = str(identificador or '').lower().strip()
    if any(k in t for k in ('encanta', 'love', '1678524932434102')):
        return 'Me encanta'
    if any(k in t for k in ('importa', 'care', 'support', '613557422527858')):
        return 'Me importa'
    if any(k in t for k in ('divierte', 'haha', '115940658764963')):
        return 'Me divierte'
    if any(k in t for k in ('asombra', 'wow', '478547315650144')):
        return 'Me asombra'
    if any(k in t for k in ('entristece', 'sad', 'sorry', '908563459236466')):
        return 'Me entristece'
    if any(k in t for k in ('enoja', 'angry', 'anger', '444813342392137')):
        return 'Me enoja'
    if any(k in t for k in ('gusta', 'like', '1635855486666999')):
        return 'Me gusta'
    return ''

def extraer_json_limpio(texto_respuesta):
    texto = texto_respuesta.strip()
    if texto.startswith("```"):
        texto = re.sub(r'^```(?:json)?\s*', '', texto)
        texto = re.sub(r'\s*```$', '', texto)
    inicio = texto.find('{')
    fin = texto.rfind('}')
    if inicio != -1 and fin != -1:
        texto = texto[inicio:fin + 1]
    return json.loads(texto)

def llamar_endpoint_gemini(prompt_texto):
    payload = {
        "contents": [{"role": "user", "parts": [{"text": prompt_texto}]}],
        "generationConfig": {
            "responseMimeType": "application/json",
            "temperature": 0.1
        }
    }

    proto = "https://"
    host_vertex = "aiplatform.googleapis.com"
    host_studio = "generativelanguage.googleapis.com"

    if GEMINI_API_KEY.startswith("AQ."):
        endpoints = [
            f"{proto}{host_vertex}/v1/publishers/google/models/gemini-2.5-flash:generateContent?key={GEMINI_API_KEY}",
            f"{proto}{host_vertex}/v1beta1/publishers/google/models/gemini-2.5-flash:generateContent?key={GEMINI_API_KEY}",
            f"{proto}{host_studio}/v1beta/models/gemini-2.5-flash:generateContent?key={GEMINI_API_KEY}"
        ]
    else:
        endpoints = [
            f"{proto}{host_studio}/v1beta/models/gemini-2.5-flash:generateContent?key={GEMINI_API_KEY}",
            f"{proto}{host_vertex}/v1/publishers/google/models/gemini-2.5-flash:generateContent?key={GEMINI_API_KEY}"
        ]

    ultimo_error = ""
    for url_api in endpoints:
        # Limpiar cualquier corchete accidental si se pegó con formato markdown
        url_limpia = url_api.replace("[", "").replace("]", "")
        if "(" in url_limpia and ")" in url_limpia:
            url_limpia = re.sub(r'\(https?://[^)]+\)', '', url_limpia)

        try:
            resp = requests.post(url_limpia, json=payload, timeout=40)
            if resp.status_code == 200:
                data_resp = resp.json()
                texto_out = data_resp['candidates'][0]['content']['parts'][0]['text']
                print("[Motor IA] Análisis de Gemini completado con éxito.")
                return extraer_json_limpio(texto_out)
            ultimo_error = f"HTTP {resp.status_code}: {resp.text[:200]}"
        except Exception as e:
            ultimo_error = str(e)

    raise RuntimeError(f"No se pudo conectar con Gemini IA ({ultimo_error})")

def analizar_todo_con_gemini(contexto_post, descripciones_alt, comentarios_crudos, usuarios_reaccion_crudos, desglose_reac, total_likes):
    actualizar_progreso(
        76,
        f"Evaluando contexto, {len(comentarios_crudos)} comentarios y {len(usuarios_reaccion_crudos)} reacciones con Gemini IA...",
        techo_etapa=93,
        segundos_holgados=6
    )

    comentarios_envio = [
        {"id": i, "username": c['username'], "comment": c['comment']}
        for i, c in enumerate(comentarios_crudos)
    ]

    tipos_reac_unicos = sorted(list(
        {r['reaccion'] for r in usuarios_reaccion_crudos if r.get('reaccion')}
        | set(desglose_reac.keys())
    ))

    prompt = (
        "Eres un analista experto de Inteligencia Artificial especializado en redes sociales en México.\n"
        "Analiza la siguiente publicación de Facebook junto con sus elementos visuales, sus comentarios y los tipos de reacción registrados.\n\n"
        f"TEXTO DE LA PUBLICACIÓN:\n{contexto_post}\n\n"
        f"ELEMENTOS VISUALES DETECTADOS:\n{json.dumps(descripciones_alt, ensure_ascii=False)}\n\n"
        f"CONTEO GLOBAL POR TIPO DE REACCIÓN:\n{json.dumps(desglose_reac, ensure_ascii=False)} (Total reacciones: {total_likes})\n\n"
        f"TIPOS DE REACCIÓN PRESENTES:\n{json.dumps(tipos_reac_unicos, ensure_ascii=False)}\n\n"
        f"COMENTARIOS EXTRAÍDOS ({len(comentarios_envio)}):\n{json.dumps(comentarios_envio, ensure_ascii=False)}\n\n"
        "TAREAS OBLIGATORIAS:\n"
        "1. Para cada comentario en COMENTARIOS EXTRAÍDOS, determina su sentimiento real ('Positivo', 'Negativo' o 'Neutro') basándote en el contexto completo de la publicación, detectando sarcasmo, ironía, burlas, quejas o apoyo genuino.\n"
        "2. Marca 'es_bot': true únicamente si hay comentarios con el mismo mensaje repetido/clonado en la lista, poniendo en 'motivo_bot' la razón (o 'Orgánico' si no es repetido).\n"
        "3. En 'evaluacion_tipos_reaccion', indica para cada tipo de reacción presente ('Me gusta', 'Me encanta', 'Me divierte', 'Me asombra', 'Me entristece', 'Me enoja', 'Me importa') qué sentimiento ('Positivo', 'Negativo' o 'Neutro') representa en el contexto específico de esta publicación (por ejemplo, si 'Me divierte' es burla hacia el post, es 'Negativo') y una breve 'explicacion'.\n"
        "4. Calcula los porcentajes globales de comentarios ('porcentajes_comentarios') y de reacciones ('reacciones_sentimiento').\n"
        "5. Extrae hasta 6 palabras clave frecuentes de los comentarios en 'palabras_frecuentes' y redacta 'contexto_multimedia_limpio' y 'resumen_ia'.\n\n"
        "Devuelve ÚNICAMENTE un JSON válido con esta estructura exacta:\n"
        "{\n"
        '  "contexto_multimedia_limpio": "descripción breve",\n'
        '  "porcentajes_comentarios": {"positivo": 0, "negativo": 0, "neutro": 0},\n'
        '  "reacciones_sentimiento": {"positivo": 0, "negativo": 0, "neutro": 0, "me_divierte_como_burla": false},\n'
        '  "total_bots": 0,\n'
        '  "resumen_ia": "resumen ejecutivo claro",\n'
        '  "palabras_frecuentes": [{"palabra": "ejemplo", "frecuencia": 3}],\n'
        '  "comentarios_evaluados": [\n'
        '    {"id": 0, "sentiment": "Positivo" | "Negativo" | "Neutro", "es_bot": false, "motivo_bot": "Orgánico"}\n'
        '  ],\n'
        '  "evaluacion_tipos_reaccion": {\n'
        '    "Me gusta": {"sentiment": "Positivo", "explicacion": "razón contextual"},\n'
        '    "Me encanta": {"sentiment": "Positivo", "explicacion": "razón contextual"},\n'
        '    "Me divierte": {"sentiment": "Positivo" | "Negativo", "explicacion": "razón contextual"},\n'
        '    "Me asombra": {"sentiment": "Neutro", "explicacion": "razón contextual"},\n'
        '    "Me entristece": {"sentiment": "Negativo", "explicacion": "razón contextual"},\n'
        '    "Me enoja": {"sentiment": "Negativo", "explicacion": "razón contextual"},\n'
        '    "Me importa": {"sentiment": "Positivo", "explicacion": "razón contextual"}\n'
        '  }\n'
        "}"
    )

    return llamar_endpoint_gemini(prompt)

def parsear_numero_fb(texto_num):
    if not texto_num:
        return 0
    texto_num = str(texto_num).lower().strip()
    multiplicador = 1
    if 'mil' in texto_num or 'k' in texto_num:
        multiplicador = 1000
        texto_num = re.sub(r'(mil|k)', '', texto_num).strip()
    elif 'm' in texto_num:
        multiplicador = 1000000
        texto_num = texto_num.replace('m', '').strip()

    texto_num = texto_num.replace(',', '.')
    match = re.search(r'\d+(\.\d+)?', texto_num)
    if not match:
        return 0
    return int(float(match.group(0)) * multiplicador)

def buscar_uri_avatar_recursivo(nodo_dict):
    if isinstance(nodo_dict, dict):
        uri = nodo_dict.get('uri')
        if isinstance(uri, str) and uri.startswith('http'):
            return uri
        for k in ('profile_picture', 'profile_picture_depth_0', 'profile_picture_depth_1', 'avatar', 'image', 'picture'):
            if k in nodo_dict:
                res = buscar_uri_avatar_recursivo(nodo_dict[k])
                if res:
                    return res
        for v in nodo_dict.values():
            if isinstance(v, dict):
                res = buscar_uri_avatar_recursivo(v)
                if res:
                    return res
    return ''

def registrar_comentario_graphql(acumulado, control_dedup, comment_id, autor, texto_com, avatar_url=''):
    """
    Registra comentarios exclusivamente desde los nodos estructurados de GraphQL.
    Usa llave estricta (ID de comentario y huella normalizada de autor + texto) para evitar cualquier duplicado.
    """
    autor_limpio = re.sub(r'\s+', ' ', (autor or '').strip())
    texto_limpio = re.sub(r'\s+', ' ', (texto_com or '').strip())
    if not autor_limpio or not texto_limpio:
        return

    # Evitar meter la descripción del post como comentario
    if acumulado['contexto_post'] and texto_limpio.lower() == re.sub(r'\s+', ' ', acumulado['contexto_post'].strip()).lower():
        return

    if comment_id and comment_id in control_dedup['ids']:
        return

    huella = f"{autor_limpio.lower()}::{re.sub(r'[^\wáéíóúñü]', '', texto_limpio.lower())}"
    if huella in control_dedup['huellas']:
        idx = control_dedup['huellas'][huella]
        if avatar_url and not acumulado['comentarios'][idx].get('avatar_url'):
            acumulado['comentarios'][idx]['avatar_url'] = avatar_url
        return

    if comment_id:
        control_dedup['ids'].add(comment_id)
    control_dedup['huellas'][huella] = len(acumulado['comentarios'])

    acumulado['comentarios'].append({
        'username': autor_limpio[:95],
        'avatar_url': avatar_url,
        'comment': texto_limpio
    })

def registrar_reactor_unico(acumulado, usuarios_reac_vistos, username, avatar_url, reaccion_raw):
    nom = re.sub(r'\s+', ' ', (username or '').strip())
    reac_limpia = traducir_nombre_reaccion(reaccion_raw)
    if not nom or not reac_limpia or len(nom) < 2:
        return

    clave = nom.lower()
    if clave in usuarios_reac_vistos:
        idx = usuarios_reac_vistos[clave]
        if avatar_url and not acumulado['usuarios_reaccion'][idx].get('avatar_url'):
            acumulado['usuarios_reaccion'][idx]['avatar_url'] = avatar_url
        # Si antes estaba como 'Me gusta' genérico y ahora llega su reacción específica de pestaña, actualizarla
        if reac_limpia != 'Me gusta' and acumulado['usuarios_reaccion'][idx].get('reaccion') == 'Me gusta':
            acumulado['usuarios_reaccion'][idx]['reaccion'] = reac_limpia
        return

    usuarios_reac_vistos[clave] = len(acumulado['usuarios_reaccion'])
    acumulado['usuarios_reaccion'].append({
        'username': nom[:95],
        'avatar_url': avatar_url or '',
        'reaccion': reac_limpia
    })

def recorrer_json_fb(nodo, acumulado, control_dedup, usuarios_reac_vistos, reaccion_pestana_activa=''):
    if isinstance(nodo, dict):
        if not acumulado['contexto_post']:
            msg = nodo.get('message')
            if isinstance(msg, dict) and isinstance(msg.get('text'), str) and len(msg['text'].strip()) > 15:
                if nodo.get('__typename') in ('Story', 'Post', '') and 'comment' not in str(nodo.get('id', '')).lower():
                    acumulado['contexto_post'] = msg['text'].strip()

        typename = str(nodo.get('__typename', ''))
        node_id = str(nodo.get('id') or nodo.get('legacy_fbid') or '')

        acc_cap = nodo.get('accessibility_caption')
        if isinstance(acc_cap, str) and len(acc_cap.strip()) > 8:
            if acc_cap.strip() not in acumulado['descripciones_alt']:
                acumulado['descripciones_alt'].append(acc_cap.strip())

        for clave_likes in ('reaction_count', 'reactors', 'likers', 'unified_reactors'):
            val = nodo.get(clave_likes)
            if isinstance(val, dict) and isinstance(val.get('count'), int) and 'node' not in nodo:
                acumulado['likes'] = max(acumulado['likes'], val['count'])

        if isinstance(nodo.get('i18n_reaction_count'), str) and 'node' not in nodo:
            acumulado['likes'] = max(acumulado['likes'], parsear_numero_fb(nodo['i18n_reaction_count']))

        # Captura de usuarios que reaccionaron desde bordes GraphQL (reactors / feedback_reaction_info)
        fb_reac_info = (
            nodo.get('feedback_reaction_info')
            or nodo.get('reaction_info')
            or nodo.get('feedback_reaction')
            or nodo.get('reaction')
        )
        sub_node = nodo.get('node')
        if isinstance(sub_node, dict) and isinstance(sub_node.get('name'), str):
            ident_reac = ''
            if isinstance(fb_reac_info, dict):
                ident_reac = (
                    fb_reac_info.get('localized_name')
                    or fb_reac_info.get('reaction_type')
                    or fb_reac_info.get('name')
                    or fb_reac_info.get('id')
                    or ''
                )
            elif reaccion_pestana_activa:
                ident_reac = reaccion_pestana_activa

            if ident_reac:
                registrar_reactor_unico(
                    acumulado,
                    usuarios_reac_vistos,
                    sub_node['name'],
                    buscar_uri_avatar_recursivo(sub_node),
                    ident_reac
                )

        # Conteo real de cada tipo de reacción en top_reactions
        top_reac = nodo.get('top_reactions')
        if isinstance(top_reac, dict) and isinstance(top_reac.get('edges'), list):
            for edge in top_reac['edges']:
                if isinstance(edge, dict):
                    cnt = edge.get('reaction_count')
                    if not isinstance(cnt, int):
                        cnt = parsear_numero_fb(edge.get('i18n_reaction_count'))
                    r_node = edge.get('node', {})
                    if isinstance(r_node, dict) and isinstance(cnt, int) and cnt > 0:
                        ident = (
                            r_node.get('localized_name')
                            or r_node.get('reaction_type')
                            or r_node.get('name')
                            or r_node.get('id')
                            or ''
                        )
                        nombre_trad = traducir_nombre_reaccion(ident)
                        if nombre_trad:
                            acumulado['reacciones'][nombre_trad] = max(acumulado['reacciones'].get(nombre_trad, 0), cnt)

        for clave_shares in ('share_count', 'reshares'):
            val = nodo.get(clave_shares)
            if isinstance(val, dict) and isinstance(val.get('count'), int):
                acumulado['compartidos'] = max(acumulado['compartidos'], val['count'])

        if isinstance(nodo.get('i18n_share_count'), str):
            acumulado['compartidos'] = max(acumulado['compartidos'], parsear_numero_fb(nodo['i18n_share_count']))

        # ÚNICAMENTE nodos estructurados de comentarios de GraphQL
        if typename == 'Comment' or (isinstance(nodo.get('body'), dict) and isinstance(nodo.get('author'), dict) and 'comment' in node_id.lower()):
            author = nodo.get('author')
            body = nodo.get('body')
            texto_com = (body.get('text') or '').strip()
            if texto_com:
                registrar_comentario_graphql(
                    acumulado,
                    control_dedup,
                    node_id,
                    author.get('name') or '',
                    texto_com,
                    buscar_uri_avatar_recursivo(author)
                )

        for valor in nodo.values():
            recorrer_json_fb(valor, acumulado, control_dedup, usuarios_reac_vistos, reaccion_pestana_activa)

    elif isinstance(nodo, list):
        for item in nodo:
            recorrer_json_fb(item, acumulado, control_dedup, usuarios_reac_vistos, reaccion_pestana_activa)

def extraer_reacciones_reales_modal(page, acumulado, usuarios_reac_vistos, estado_tab):
    """
    Abre el diálogo de reacciones de la publicación, identifica cada pestaña de reacción
    (Me gusta, Me encanta, Me divierte, Me asombra, Me entristece, Me enoja, Me importa)
    y extrae los usuarios reales de cada pestaña junto con su foto de perfil.
    """
    try:
        abrio = page.evaluate("""() => {
            const candidatos = Array.from(document.querySelectorAll('[aria-label], div[role="button"]'));
            for (const el of candidatos) {
                const aria = (el.getAttribute('aria-label') || '').trim();
                const txt = (el.innerText || '').trim();
                if (/^(Me gusta|Me encanta|Me importa|Me divierte|Me asombra|Me entristece|Me enoja|Like|Love|Care|Haha|Wow|Sad|Angry):\\s*\\d/i.test(aria) ||
                    /quién reaccionó|who reacted|todas las reacciones|ver quién reaccionó/i.test(aria) ||
                    (/^todas las reacciones/i.test(txt) && txt.length < 50)) {
                    el.click();
                    return true;
                }
            }
            return false;
        }""")
        if not abrio:
            return

        page.wait_for_timeout(1200)

        # Buscar el diálogo que contiene las pestañas de reacciones (div[role="tab"]) y no contiene artículos de comentarios
        tabs_info = page.evaluate("""() => {
            const dialogs = Array.from(document.querySelectorAll('div[role="dialog"]'));
            const reacDialog = dialogs.find(d => d.querySelectorAll('div[role="tab"]').length > 0 && d.querySelectorAll('div[role="article"]').length === 0);
            if (!reacDialog) return null;

            return Array.from(reacDialog.querySelectorAll('div[role="tab"]')).map((t, idx) => {
                const aria = t.getAttribute('aria-label') || '';
                const subAria = Array.from(t.querySelectorAll('[aria-label]')).map(x => x.getAttribute('aria-label')).join(' ');
                const imgAlt = Array.from(t.querySelectorAll('img')).map(x => x.getAttribute('alt') || '').join(' ');
                return {idx, rawLabel: `${aria} ${subAria} ${imgAlt} ${t.innerText || ''}`.trim()};
            });
        }""")

        if not tabs_info:
            return

        # Recorrer primero las pestañas específicas de cada reacción (saltando la pestaña 0 "Todas")
        for tab in tabs_info:
            reac_nombre = traducir_nombre_reaccion(tab['rawLabel'])
            if not reac_nombre:
                continue

            cant_tab = parsear_numero_fb(tab['rawLabel'])
            if cant_tab > 0:
                acumulado['reacciones'][reac_nombre] = max(acumulado['reacciones'].get(reac_nombre, 0), cant_tab)

            estado_tab['reaccion'] = reac_nombre
            try:
                page.evaluate("""(idx) => {
                    const dialogs = Array.from(document.querySelectorAll('div[role="dialog"]'));
                    const reacDialog = dialogs.find(d => d.querySelectorAll('div[role="tab"]').length > 0 && d.querySelectorAll('div[role="article"]').length === 0);
                    if (!reacDialog) return;
                    const tabs = reacDialog.querySelectorAll('div[role="tab"]');
                    if (tabs[idx]) tabs[idx].click();
                }""", tab['idx'])
                page.wait_for_timeout(700)

                # Hacer scroll dentro de esa pestaña para cargar más personas de esa reacción
                filas = page.evaluate("""() => {
                    const dialogs = Array.from(document.querySelectorAll('div[role="dialog"]'));
                    const reacDialog = dialogs.find(d => d.querySelectorAll('div[role="tab"]').length > 0 && d.querySelectorAll('div[role="article"]').length === 0);
                    if (!reacDialog) return [];

                    const divs = reacDialog.querySelectorAll('div');
                    for (const d of divs) {
                        if (d.scrollHeight > d.clientHeight + 40) {
                            d.scrollTop = d.scrollHeight;
                        }
                    }

                    const extraidos = [];
                    const svgs = reacDialog.querySelectorAll('svg');
                    for (const svg of svgs) {
                        const imgEl = svg.querySelector('image');
                        if (!imgEl) continue;
                        const avatarUrl = imgEl.getAttribute('xlink:href') || imgEl.getAttribute('href') || '';
                        if (!avatarUrl || !avatarUrl.startsWith('http')) continue;

                        let fila = svg.parentElement;
                        for (let n = 0; n < 6 && fila && fila !== reacDialog; n++) {
                            const enlaces = Array.from(fila.querySelectorAll('a[role="link"]'));
                            const TextosValidos = enlaces
                                .map(a => (a.innerText || '').trim())
                                .filter(t => t.length >= 3 && !t.includes('\\n'));
                            if (TextosValidos.length > 0) {
                                extraidos.push({
                                    username: TextosValidos[0],
                                    avatar_url: avatarUrl
                                });
                                break;
                            }
                            fila = fila.parentElement;
                        }
                    }
                    return extraidos;
                }""")

                for f in filas:
                    registrar_reactor_unico(
                        acumulado,
                        usuarios_reac_vistos,
                        f.get('username', ''),
                        f.get('avatar_url', ''),
                        reac_nombre
                    )
            except Exception:
                continue

        estado_tab['reaccion'] = ''
        page.keyboard.press('Escape')
        page.wait_for_timeout(350)
    except Exception as e:
        print(f"Aviso en modal de reacciones: {e}")

def _tarea_scraping_en_navegador_caliente(url):
    acumulado = {
        'likes': 0,
        'compartidos': 0,
        'contexto_post': '',
        'descripciones_alt': [],
        'comentarios': [],
        'usuarios_reaccion': [],
        'reacciones': {}
    }
    control_dedup = {'ids': set(), 'huellas': {}}
    usuarios_reac_vistos = {}
    estado_tab = {'reaccion': ''}

    if NAVEGADOR_RAM['context'] is None:
        _iniciar_navegador_invisible_worker()

    context = NAVEGADOR_RAM['context']
    page = context.new_page()

    def interceptar_graphql(response):
        if 'graphql' in response.url and response.status == 200:
            try:
                for linea in response.text().splitlines():
                    linea = linea.strip()
                    if linea.startswith('{'):
                        recorrer_json_fb(
                            json.loads(linea),
                            acumulado,
                            control_dedup,
                            usuarios_reac_vistos,
                            estado_tab['reaccion']
                        )
            except Exception:
                pass

    page.on("response", interceptar_graphql)

    try:
        actualizar_progreso(15, "Abriendo publicación en Facebook...", techo_etapa=26, segundos_holgados=18)
        page.goto(url, wait_until='domcontentloaded', timeout=15000)

        try:
            page.wait_for_selector('script[type="application/json"]', state='attached', timeout=1800)
        except Exception:
            page.wait_for_timeout(600)

        try:
            og_title = page.locator('meta[property="og:title"]').first.get_attribute('content') or ''
            og_desc = page.locator('meta[property="og:description"]').first.get_attribute('content') or ''
            texto_meta = og_desc if len(og_desc) > len(og_title) else og_title
            if texto_meta:
                acumulado['contexto_post'] = texto_meta.strip()
        except Exception:
            pass

        scripts_json = page.locator('script[type="application/json"]').all_inner_texts()
        for raw_json in scripts_json:
            try:
                recorrer_json_fb(json.loads(raw_json), acumulado, control_dedup, usuarios_reac_vistos)
            except Exception:
                continue

        # 1. Extraer reacciones por pestaña (Me encanta, Me divierte, Me asombra, Me enoja, Me gusta...)
        actualizar_progreso(26, "Extrayendo usuarios por cada tipo de reacción...", techo_etapa=38, segundos_holgados=14)
        extraer_reacciones_reales_modal(page, acumulado, usuarios_reac_vistos, estado_tab)

        # 2. Cambiar el selector de comentarios a "Todos los comentarios"
        try:
            abrio_menu = page.evaluate("""() => {
                const spans = Array.from(document.querySelectorAll('div[role="button"] span, span'));
                for (const s of spans) {
                    const t = (s.innerText || '').trim().toLowerCase();
                    if (t === 'más relevantes' || t === 'most relevant' || t === 'más recientes') {
                        (s.closest('div[role="button"]') || s).click();
                        return true;
                    }
                }
                return false;
            }""")
            if abrio_menu:
                page.wait_for_timeout(550)
                page.evaluate("""() => {
                    const items = Array.from(document.querySelectorAll('div[role="menuitem"], div[role="menuitemradio"], span'));
                    for (const el of items) {
                        const t = (el.innerText || '').trim().toLowerCase();
                        if (t.startsWith('todos los comentarios') || t.startsWith('all comments')) {
                            el.click();
                            return true;
                        }
                    }
                    return false;
                }""")
                page.wait_for_timeout(1000)
        except Exception:
            pass

        # 3. Desplegar comentarios paginados (los datos limpios entran automáticamente por interceptar_graphql sin duplicarse)
        sin_cambios = 0
        for ciclo in range(22):
            antes = len(acumulado['comentarios'])

            clics_realizados = page.evaluate("""() => {
                let clics = 0;
                const botones = Array.from(document.querySelectorAll('div[role="button"], span[dir="auto"]'));
                for (const b of botones) {
                    const t = (b.innerText || '').trim().toLowerCase();
                    if (t.includes('ver más comentarios') || t.includes('ver comentarios anteriores') ||
                        t.includes('view more comments') || /^ver\\s+\\d+\\s+comentario/.test(t) ||
                        /^\\d+\\s+comentario.*más/.test(t) || /^ver\\s+\\d+\\s+respuesta/.test(t) ||
                        /^\\d+\\s+respuesta/.test(t) || t === 'ver respuesta' ||
                        /^ver\\s+todas\\s+las\\s+\\d+\\s+respuestas/.test(t)) {
                        b.click();
                        clics++;
                    }
                }

                // Desplazar exactamente el contenedor con scroll donde viven los comentarios
                const articulos = document.querySelectorAll('div[role="article"]');
                if (articulos.length > 0) {
                    const ultimo = articulos[articulos.length - 1];
                    ultimo.scrollIntoView({behavior: 'instant', block: 'end'});
                    let p = ultimo.parentElement;
                    while (p) {
                        if (p.scrollHeight > p.clientHeight + 20) {
                            p.scrollTop = p.scrollHeight;
                        }
                        p = p.parentElement;
                    }
                }
                window.scrollBy(0, 1600);
                return clics;
            }""")

            page.wait_for_timeout(700 if clics_realizados > 0 else 450)

            actuales = len(acumulado['comentarios'])
            porc_scroll = min(73, 38 + (ciclo + 1) * 2)
            actualizar_progreso(
                porc_scroll,
                f"Desplegando comentarios de la publicación ({actuales} extraídos)...",
                techo_etapa=75,
                segundos_holgados=max(5, 16 - ciclo)
            )

            if actuales == antes and clics_realizados == 0:
                sin_cambios += 1
                if sin_cambios >= 3:
                    break
            else:
                sin_cambios = 0

        suma_desglose = sum(acumulado['reacciones'].values())
        if suma_desglose > acumulado['likes']:
            acumulado['likes'] = suma_desglose

        print(
            f"[Scraper FB] Likes: {acumulado['likes']} | Desglose: {acumulado['reacciones']} | "
            f"Usuarios con reacción: {len(acumulado['usuarios_reaccion'])} | "
            f"Comentarios únicos (GraphQL): {len(acumulado['comentarios'])}"
        )

    except Exception as e:
        print(f"Aviso en extracción: {e}")
    finally:
        page.close()

    return (
        acumulado['likes'],
        acumulado['compartidos'],
        acumulado['comentarios'],
        acumulado['usuarios_reaccion'],
        acumulado['reacciones'],
        acumulado['contexto_post'],
        acumulado['descripciones_alt']
    )

def extraer_datos_facebook(url):
    futuro = SCRAPER_EXECUTOR.submit(_tarea_scraping_en_navegador_caliente, url)
    return futuro.result()

@app.route('/api/avatar', methods=['GET'])
def proxy_avatar():
    url_img = request.args.get('url', '').strip()
    if not url_img or not url_img.startswith('http'):
        return '', 404
    try:
        resp = requests.get(
            url_img,
            headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'},
            timeout=4
        )
        if resp.status_code == 200:
            return Response(resp.content, mimetype=resp.headers.get('Content-Type', 'image/jpeg'))
    except Exception:
        pass
    return '', 404

@app.route('/api/progreso', methods=['GET'])
def consultar_progreso():
    if ESTADO_PROGRESO['activo'] and ESTADO_PROGRESO['inicio_ts'] > 0:
        ahora = time.time()
        ESTADO_PROGRESO['tiempo_transcurrido'] = max(0, int(ahora - ESTADO_PROGRESO['inicio_ts']))

        if ESTADO_PROGRESO['porcentaje'] < ESTADO_PROGRESO['techo_etapa']:
            ESTADO_PROGRESO['porcentaje'] += 2

        restante_calculado = max(3, int(round(ESTADO_PROGRESO['meta_holgada_ts'] - ahora)))
        ESTADO_PROGRESO['tiempo_restante'] = min(ESTADO_PROGRESO['tiempo_restante'], restante_calculado)

    return jsonify(ESTADO_PROGRESO), 200

@app.route('/api/login', methods=['POST'])
def login():
    data = request.get_json() or {}
    correo = (data.get('correo') or '').strip()
    password = data.get('password') or ''

    if not correo or not password:
        return jsonify({'error': 'Ingresa tu correo y contraseña'}), 400

    try:
        conn = get_db_connection()
        with conn.cursor() as cursor:
            cursor.execute("SELECT * FROM usuarios WHERE correo = %s", (correo,))
            user = cursor.fetchone()
        conn.close()

        if not user:
            return jsonify({'error': 'No existe una cuenta con ese correo'}), 401

        stored_pw = user['password']
        es_valido = check_password_hash(stored_pw, password) if stored_pw.startswith(('scrypt:', 'pbkdf2:')) else (stored_pw == password)

        if es_valido:
            return jsonify({
                'mensaje': 'Acceso correcto',
                'usuario': {
                    'id': user['id'],
                    'nombre': user['nombre'],
                    'correo': user['correo'],
                    'rol': user['rol']
                }
            }), 200

        return jsonify({'error': 'La contraseña ingresada es incorrecta'}), 401
    except Exception as e:
        return jsonify({'error': f'Error en el servidor: {str(e)}'}), 500

@app.route('/api/registro', methods=['POST'])
def registro():
    data = request.get_json() or {}
    nombre = (data.get('nombre') or '').strip()
    correo = (data.get('correo') or '').strip()
    password = data.get('password') or ''
    rol = 'analista'

    if not nombre or not correo or len(password) < 4:
        return jsonify({'error': 'Completa todos los campos (contraseña de al menos 4 caracteres)'}), 400

    try:
        conn = get_db_connection()
        with conn.cursor() as cursor:
            cursor.execute("SELECT id FROM usuarios WHERE correo = %s", (correo,))
            if cursor.fetchone():
                conn.close()
                return jsonify({'error': 'Este correo electrónico ya está registrado'}), 409

            hash_pw = generate_password_hash(password)
            cursor.execute(
                "INSERT INTO usuarios (nombre, correo, password, rol) VALUES (%s, %s, %s, %s)",
                (nombre, correo, hash_pw, rol)
            )
            nuevo_id = cursor.lastrowid
            conn.commit()
        conn.close()

        return jsonify({
            'mensaje': 'Usuario registrado correctamente',
            'usuario': {'id': nuevo_id, 'nombre': nombre, 'correo': correo, 'rol': rol}
        }), 201
    except Exception as e:
        return jsonify({'error': f'Error al registrar usuario: {str(e)}'}), 500

@app.route('/api/analizar', methods=['POST'])
def analizar():
    data = request.get_json() or {}
    url = data.get('url', '').strip()

    if not url or ('facebook.com' not in url and 'fb.watch' not in url and 'fb.com' not in url):
        return jsonify({'error': 'Ingresa un enlace válido de Facebook'}), 400

    iniciar_medidor_progreso(colchon_segundos=25)

    try:
        (
            likes,
            compartidos,
            comentarios_raw,
            usuarios_reaccion_raw,
            desglose_reac,
            contexto_post,
            descripciones_alt
        ) = extraer_datos_facebook(url)

        if likes == 0 and compartidos == 0 and len(comentarios_raw) == 0:
            actualizar_progreso(100, "Extracción detenida por restricción de Facebook.")
            return jsonify({
                'error': 'Facebook bloqueó temporalmente la lectura de este enlace. Intenta de nuevo en unos segundos.'
            }), 429

        resultado_ia = analizar_todo_con_gemini(
            contexto_post,
            descripciones_alt,
            comentarios_raw,
            usuarios_reaccion_raw,
            desglose_reac,
            likes
        )

        mapa_eval_com = {
            int(item['id']): item
            for item in resultado_ia.get('comentarios_evaluados', [])
            if 'id' in item
        }

        comentarios_procesados = []
        for idx, c in enumerate(comentarios_raw):
            ev = mapa_eval_com.get(idx, {})
            comentarios_procesados.append({
                'username': c['username'],
                'avatar_url': c.get('avatar_url', ''),
                'comment': c['comment'],
                'sentiment': ev.get('sentiment', 'Neutro'),
                'es_bot': bool(ev.get('es_bot', False)),
                'motivo_bot': ev.get('motivo_bot', 'Orgánico'),
                'confianza_ia': 99.0
            })

        eval_tipos_reac = resultado_ia.get('evaluacion_tipos_reaccion') or {}
        lista_reacciones = []

        for u in usuarios_reaccion_raw:
            nom_r = u['reaccion']
            info_r = eval_tipos_reac.get(nom_r, {})
            lista_reacciones.append({
                'username': u['username'],
                'avatar_url': u.get('avatar_url', ''),
                'reaccion': nom_r,
                'sentiment': info_r.get('sentiment', 'Positivo'),
                'explicacion': info_r.get('explicacion', 'Evaluada según el contexto de la publicación.')
            })

        # Si Facebook no abrió la lista de nombres en ese post pero sí dio los conteos por tipo de reacción
        if not lista_reacciones and desglose_reac:
            for nom_r, cant in desglose_reac.items():
                if cant > 0:
                    info_r = eval_tipos_reac.get(nom_r, {})
                    lista_reacciones.append({
                        'username': f'{cant} personas',
                        'avatar_url': '',
                        'reaccion': nom_r,
                        'sentiment': info_r.get('sentiment', 'Positivo'),
                        'explicacion': info_r.get('explicacion', 'Evaluada según el contexto de la publicación.')
                    })

        porc_com = resultado_ia.get('porcentajes_comentarios') or {}
        porc_pos = int(porc_com.get('positivo', 0))
        porc_neg = int(porc_com.get('negativo', 0))
        porc_neu = int(porc_com.get('neutro', 0))

        reac_sent = resultado_ia.get('reacciones_sentimiento') or {
            'positivo': 100, 'negativo': 0, 'neutro': 0, 'me_divierte_como_burla': False
        }
        total_bots = sum(1 for c in comentarios_procesados if c['es_bot'])
        resumen_ia = resultado_ia.get('resumen_ia') or ''
        contexto_multimedia = resultado_ia.get('contexto_multimedia_limpio') or 'Publicación analizada.'
        palabras_frecuentes = resultado_ia.get('palabras_frecuentes') or []
        total_com = len(comentarios_procesados)

        actualizar_progreso(95, "Guardando resultados en MariaDB...", techo_etapa=99, segundos_holgados=2)

        texto_post_final = contexto_post or 'Publicación analizada con IA.'
        reacciones_json_str = json.dumps(lista_reacciones, ensure_ascii=False)

        conn = get_db_connection()
        with conn.cursor() as cursor:
            cursor.execute("DELETE FROM analisis WHERE url_post = %s", (url[:495],))

            sql_analisis = """
                INSERT INTO analisis (
                    url_post, total_likes, total_compartidos, total_comentarios,
                    positivos, negativos, neutros, resumen_ia,
                    contexto_post, contexto_multimedia, reacciones_lista_json, total_bots
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """
            cursor.execute(sql_analisis, (
                url[:495], likes, compartidos, total_com,
                porc_pos, porc_neg, porc_neu, resumen_ia,
                texto_post_final, contexto_multimedia, reacciones_json_str, total_bots
            ))
            id_analisis = cursor.lastrowid

            if comentarios_procesados:
                sql_com = """
                    INSERT INTO comentarios (
                        id_analisis, usuario, comentario, sentimiento,
                        es_bot, motivo_bot, avatar_url, confianza_ia
                    )
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                """
                datos_lote = [
                    (
                        id_analisis,
                        c['username'],
                        c['comment'],
                        c['sentiment'],
                        1 if c['es_bot'] else 0,
                        c['motivo_bot'],
                        c.get('avatar_url', ''),
                        c.get('confianza_ia', 99.0)
                    )
                    for c in comentarios_procesados
                ]
                cursor.executemany(sql_com, datos_lote)

            conn.commit()
        conn.close()

        actualizar_progreso(100, "Análisis finalizado correctamente.")

        return jsonify({
            'id_analisis': id_analisis,
            'url_post': url,
            'contexto_post': texto_post_final,
            'contexto_multimedia': contexto_multimedia,
            'likes': likes,
            'compartidos': compartidos,
            'total_comentarios': total_com,
            'total_bots': total_bots,
            'porcentajes': {
                'positivo': porc_pos,
                'negativo': porc_neg,
                'neutro': porc_neu
            },
            'reacciones_sentimiento': reac_sent,
            'reacciones_lista': lista_reacciones,
            'palabras_frecuentes': palabras_frecuentes,
            'resumen_ia': resumen_ia,
            'comentarios': comentarios_procesados
        }), 200

    except Exception as e:
        actualizar_progreso(100, "Error durante el procesamiento.")
        print(f"[Error Analizar] {e}")
        return jsonify({'error': str(e)}), 500

@app.route('/api/historial', methods=['GET'])
def obtener_historial():
    try:
        conn = get_db_connection()
        with conn.cursor() as cursor:
            cursor.execute(
                "SELECT id, url_post, total_likes, total_compartidos, total_comentarios, total_bots, "
                "positivos, negativos, neutros, fecha_analisis "
                "FROM analisis ORDER BY fecha_analisis DESC LIMIT 30"
            )
            filas = cursor.fetchall()
        conn.close()

        for f in filas:
            f['fecha_analisis'] = str(f['fecha_analisis'])
        return jsonify({'historial': filas}), 200
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/historial/<int:id_analisis>', methods=['GET'])
def cargar_analisis_guardado(id_analisis):
    try:
        conn = get_db_connection()
        with conn.cursor() as cursor:
            cursor.execute("SELECT * FROM analisis WHERE id = %s", (id_analisis,))
            analisis = cursor.fetchone()
            if not analisis:
                conn.close()
                return jsonify({'error': 'Registro no encontrado'}), 404

            cursor.execute(
                "SELECT usuario AS username, comentario AS comment, sentimiento AS sentiment, "
                "es_bot, motivo_bot, avatar_url, confianza_ia FROM comentarios WHERE id_analisis = %s",
                (id_analisis,)
            )
            comentarios = cursor.fetchall()
        conn.close()

        for c in comentarios:
            c['es_bot'] = bool(c.get('es_bot', 0))
            c['motivo_bot'] = c.get('motivo_bot') or 'Orgánico'
            c['avatar_url'] = c.get('avatar_url') or ''

        lista_reacciones = []
        if analisis.get('reacciones_lista_json'):
            try:
                lista_reacciones = json.loads(analisis['reacciones_lista_json'])
            except Exception:
                lista_reacciones = []

        return jsonify({
            'id_analisis': analisis['id'],
            'url_post': analisis['url_post'],
            'contexto_post': analisis.get('contexto_post') or 'Consulta cargada desde el historial.',
            'contexto_multimedia': analisis.get('contexto_multimedia') or 'Contexto multimedia almacenado previamente.',
            'likes': analisis['total_likes'],
            'compartidos': analisis['total_compartidos'],
            'total_comentarios': analisis['total_comentarios'],
            'total_bots': analisis.get('total_bots', 0),
            'porcentajes': {
                'positivo': analisis['positivos'],
                'negativo': analisis['negativos'],
                'neutro': analisis['neutros']
            },
            'reacciones_sentimiento': {
                'positivo': analisis['positivos'],
                'negativo': analisis['negativos'],
                'neutro': analisis['neutros'],
                'me_divierte_como_burla': analisis['negativos'] >= analisis['positivos']
            },
            'reacciones_lista': lista_reacciones,
            'palabras_frecuentes': [],
            'resumen_ia': analisis['resumen_ia'],
            'comentarios': comentarios
        }), 200
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/historial/<int:id_analisis>', methods=['DELETE'])
def eliminar_analisis(id_analisis):
    try:
        conn = get_db_connection()
        with conn.cursor() as cursor:
            cursor.execute("DELETE FROM analisis WHERE id = %s", (id_analisis,))
            conn.commit()
        conn.close()
        return jsonify({'mensaje': 'Registro eliminado'}), 200
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/exportar/<int:id_analisis>', methods=['GET'])
def exportar_excel(id_analisis):
    try:
        conn = get_db_connection()
        with conn.cursor() as cursor:
            cursor.execute("SELECT * FROM analisis WHERE id = %s", (id_analisis,))
            analisis = cursor.fetchone()
            cursor.execute(
                "SELECT usuario, comentario, sentimiento, es_bot, motivo_bot "
                "FROM comentarios WHERE id_analisis = %s",
                (id_analisis,)
            )
            comentarios = cursor.fetchall()
        conn.close()

        if not analisis:
            return jsonify({'error': 'Registro no encontrado'}), 404

        lista_reacciones = []
        if analisis.get('reacciones_lista_json'):
            try:
                lista_reacciones = json.loads(analisis['reacciones_lista_json'])
            except Exception:
                lista_reacciones = []

        wb = Workbook()
        ws_resumen = wb.active
        ws_resumen.title = "Resumen y gráficas"

        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)

        ws_resumen["A1"] = "Reporte de análisis de publicación en Facebook"
        ws_resumen["A1"].font = Font(size=14, bold=True, color="0F172A")
        ws_resumen["A2"] = "Enlace analizado:"
        ws_resumen["B2"] = analisis['url_post']
        ws_resumen["A3"] = "Fecha de consulta:"
        ws_resumen["B3"] = str(analisis['fecha_analisis'])
        ws_resumen["A4"] = "Comentarios repetidos (bots):"
        ws_resumen["B4"] = analisis.get('total_bots', 0)
        ws_resumen["A5"] = "Resumen de IA:"
        ws_resumen["B5"] = analisis['resumen_ia']

        ws_resumen["A7"] = "Sentimiento en comentarios"
        ws_resumen["B7"] = "Porcentaje"
        ws_resumen["A7"].fill, ws_resumen["B7"].fill = header_fill, header_fill
        ws_resumen["A7"].font, ws_resumen["B7"].font = header_font, header_font

        ws_resumen["A8"] = "Positivos"
        ws_resumen["B8"] = analisis['positivos']
        ws_resumen["A9"] = "Negativos"
        ws_resumen["B9"] = analisis['negativos']
        ws_resumen["A10"] = "Neutros"
        ws_resumen["B10"] = analisis['neutros']

        grafica_pastel = PieChart()
        grafica_pastel.title = "Distribución de sentimientos en comentarios (%)"
        etiquetas_ref = Reference(ws_resumen, min_col=1, min_row=8, max_row=10)
        datos_ref = Reference(ws_resumen, min_col=2, min_row=7, max_row=10)
        grafica_pastel.add_data(datos_ref, titles_from_data=True)
        grafica_pastel.set_categories(etiquetas_ref)
        grafica_pastel.width = 14
        grafica_pastel.height = 10
        ws_resumen.add_chart(grafica_pastel, "D6")

        ws_resumen["A13"] = "Usuario que reaccionó"
        ws_resumen["B13"] = "Reacción"
        ws_resumen["C13"] = "Sentimiento asignado por IA"
        for col in ("A13", "B13", "C13"):
            ws_resumen[col].fill = header_fill
            ws_resumen[col].font = header_font

        fila_actual = 14
        for r in lista_reacciones:
            ws_resumen[f"A{fila_actual}"] = r.get('username') or r.get('reaccion', '')
            ws_resumen[f"B{fila_actual}"] = r.get('reaccion', '')
            ws_resumen[f"C{fila_actual}"] = r.get('sentiment', '')
            fila_actual += 1

        ws_resumen.column_dimensions['A'].width = 30
        ws_resumen.column_dimensions['B'].width = 25
        ws_resumen.column_dimensions['C'].width = 26

        ws_com = wb.create_sheet(title="Comentarios")
        ws_com.append(["Nombre del usuario", "Comentario", "Sentimiento IA", "Clasificación"])
        for col in ("A1", "B1", "C1", "D1"):
            ws_com[col].fill = header_fill
            ws_com[col].font = header_font

        for c in comentarios:
            ws_com.append([
                c['usuario'],
                c['comentario'],
                c['sentimiento'],
                c.get('motivo_bot') or 'Orgánico'
            ])

        ws_com.column_dimensions['A'].width = 28
        ws_com.column_dimensions['B'].width = 68
        ws_com.column_dimensions['C'].width = 20
        ws_com.column_dimensions['D'].width = 28

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        return send_file(
            output,
            as_attachment=True,
            download_name=f"reporte_facebook_{id_analisis}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    except Exception as e:
        return jsonify({'error': str(e)}), 500

if __name__ == '__main__':
    port = int(os.getenv('PORT', 5000))
    app.run(debug=True, use_reloader=False, threaded=True, port=port)